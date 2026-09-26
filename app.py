from __future__ import annotations

import hashlib
import hmac
import json
import os
import re
import tempfile
import unicodedata
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

import requests
import streamlit as st
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from streamlit_autorefresh import st_autorefresh

APP_TITLE = "Consulta MDU Óptica"
TARGET_SHEET = "MDU'S BROWNFIELD(CONSULTA)"
FALLBACK_PATH = Path(__file__).parent / "data" / "base_mdu.json"
REFRESH_SECONDS = 300
PAGE_SIZE = 50
DEFAULT_SOURCE_URL = "https://drive.google.com/drive/folders/1h8Ium1WG9ZmZeuONQuBtAGScw7ukie9V?usp=drive_link"

st.set_page_config(
    page_title=APP_TITLE,
    page_icon="🔎",
    layout="wide",
    initial_sidebar_state="collapsed",
)


def secret(name: str, default: str = "") -> str:
    try:
        return str(st.secrets.get(name, default)).strip()
    except Exception:
        return str(os.getenv(name, default)).strip()


def normalize_header(value: Any) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFD", str(value))
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    text = re.sub(r"[^A-Z0-9]+", " ", text.upper())
    return re.sub(r"\s+", " ", text).strip()


def normalize_search(value: Any) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFD", str(value))
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    text = text.upper().strip()
    text = re.sub(r"\bRUA\b", "R", text)
    text = re.sub(r"\bAVENIDA\b", "AV", text)
    text = re.sub(r"[^A-Z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def clean_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if float(value) == 0:
            return ""
    value = str(value).strip()
    if not value:
        return ""
    if value.upper() in {"0", "#N/A", "#REF!", "#VALUE!", "#NAME?", "N/A", "NA", "-", "NONE"}:
        return ""
    return value


def format_date(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, (int, float)):
        number = float(value)
        if number == 0 or (2000 <= number <= 2050 and number.is_integer()):
            return ""
        if number > 1000:
            try:
                return from_excel(number).strftime("%d/%m/%Y")
            except Exception:
                pass
    text = clean_value(value)
    if not text:
        return ""
    try:
        number = float(text.replace(",", "."))
        if number > 1000 and not (2000 <= number <= 2050 and number.is_integer()):
            return from_excel(number).strftime("%d/%m/%Y")
    except Exception:
        pass
    return text


def extract_drive_id(url: str) -> str:
    patterns = [
        r"/file/d/([A-Za-z0-9_-]+)",
        r"[?&]id=([A-Za-z0-9_-]+)",
        r"/d/([A-Za-z0-9_-]+)",
    ]
    for pattern in patterns:
        match = re.search(pattern, url)
        if match:
            return match.group(1)
    return ""


def source_date_from_name(name: str) -> str:
    match = re.search(r"(?:^|[^0-9])(\d{2})[_-](\d{2})[_-](\d{2}|\d{4})(?:[^0-9]|$)", name or "")
    if not match:
        return ""
    day, month, year = match.groups()
    if len(year) == 2:
        year = f"20{year}"
    try:
        parsed = datetime(int(year), int(month), int(day))
        return parsed.strftime("%d/%m/%Y")
    except ValueError:
        return ""


def download_source(url: str) -> tuple[bytes, str]:
    headers = {
        "User-Agent": "Mozilla/5.0 Consulta-MDU/2.0",
        "Cache-Control": "no-cache, no-store, max-age=0",
        "Pragma": "no-cache",
    }

    if "drive.google.com/drive/folders/" in url:
        try:
            import gdown
            with tempfile.TemporaryDirectory(prefix="mdu_drive_") as tmpdir:
                result = gdown.download_folder(
                    url=url,
                    output=tmpdir,
                    quiet=True,
                    use_cookies=False,
                    remaining_ok=True,
                )
                if not result:
                    raise RuntimeError("A pasta do Google Drive está vazia ou não está acessível ao site.")
                candidates = []
                for item in result:
                    path = Path(item)
                    if path.suffix.lower() in {".xlsx", ".xlsm"} and not path.name.startswith("~$"):
                        candidates.append(path)
                if not candidates:
                    candidates = [p for p in Path(tmpdir).rglob("*.xlsx") if not p.name.startswith("~$")]
                if len(candidates) == 0:
                    raise RuntimeError("Nenhuma planilha .xlsx foi encontrada na pasta MDU-FORECAST.")
                if len(candidates) > 1:
                    names = ", ".join(sorted(p.name for p in candidates)[:6])
                    raise RuntimeError(f"Há mais de uma planilha na pasta ({names}). Deixe somente o FORECAST atual.")
                source = candidates[0]
                content = source.read_bytes()
                if len(content) < 100_000:
                    raise RuntimeError("A planilha encontrada parece incompleta ou inválida.")
                return content, f"Google Drive · {source.name}"
        except Exception as exc:
            raise RuntimeError(f"Falha ao ler a pasta MDU-FORECAST: {exc}") from exc

    if "drive.google.com" in url:
        try:
            import gdown

            with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
                tmp_path = tmp.name
            try:
                result = gdown.download(url=url, output=tmp_path, quiet=True, fuzzy=True, use_cookies=False)
                if not result:
                    raise RuntimeError("O Google Drive não liberou o download da planilha.")
                content = Path(tmp_path).read_bytes()
                return content, "Google Drive"
            finally:
                Path(tmp_path).unlink(missing_ok=True)
        except Exception as exc:
            drive_id = extract_drive_id(url)
            if not drive_id:
                raise RuntimeError(f"Falha ao baixar do Google Drive: {exc}") from exc
            url = f"https://drive.usercontent.google.com/download?id={drive_id}&export=download&confirm=t"

    response = requests.get(url, timeout=180, headers=headers, allow_redirects=True)
    response.raise_for_status()
    content = response.content
    if len(content) < 100_000:
        content_type = response.headers.get("content-type", "")
        if "html" in content_type.lower():
            raise RuntimeError("O link retornou uma página HTML, não a planilha. Use um link de download direto/compartilhado.")
    return content, response.url


def extract_records_from_xlsx(content: bytes, source_name: str) -> dict[str, Any]:
    wb = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
    try:
        if TARGET_SHEET not in wb.sheetnames:
            raise RuntimeError(f"A aba '{TARGET_SHEET}' não foi encontrada.")
        ws = wb[TARGET_SHEET]

        expected = {
            "city": "DSC CIDADE",
            "address": "ENDERECO MDU",
            "node": "NODE",
            "epo": "EPO MDU RESPONSAVEL",
            "finish": "CONSTRUCAO RI MDU FIM",
            "status": "CONSTRUCAO RI MDU STATUS",
        }

        header_row = None
        col_map: dict[str, int] = {}
        for row_num, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
            headers = {normalize_header(value): idx for idx, value in enumerate(row) if normalize_header(value)}
            if all(label in headers for label in expected.values()):
                header_row = row_num
                col_map = {key: headers[label] for key, label in expected.items()}
                break

        if header_row is None:
            raise RuntimeError("Não foi possível localizar os 6 campos obrigatórios na aba de consulta.")

        rows: list[dict[str, str]] = []
        seen: set[tuple[str, str, str, str, str, str]] = set()
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            city = clean_value(row[col_map["city"]] if col_map["city"] < len(row) else None)
            address = clean_value(row[col_map["address"]] if col_map["address"] < len(row) else None)
            if not city or not address:
                continue
            node = clean_value(row[col_map["node"]] if col_map["node"] < len(row) else None)
            epo = clean_value(row[col_map["epo"]] if col_map["epo"] < len(row) else None).upper()
            finish = format_date(row[col_map["finish"]] if col_map["finish"] < len(row) else None)
            status = clean_value(row[col_map["status"]] if col_map["status"] < len(row) else None)
            key = (city, address, node, epo, finish, status)
            if key in seen:
                continue
            seen.add(key)
            rows.append({"c": city, "a": address, "n": node, "e": epo, "f": finish, "s": status})

        rows.sort(key=lambda item: (item["c"], item["a"], item["n"], item["e"], item["f"], item["s"]))
        cities = sorted({item["c"] for item in rows})
        if len(rows) < 1000:
            raise RuntimeError(f"Carga bloqueada: somente {len(rows)} registros válidos foram encontrados.")

        return {
            "meta": {
                "sourceDate": source_date_from_name(source_name) or "não identificada",
                "updatedAt": datetime.now().strftime("%d/%m/%Y %H:%M"),
                "sourceFile": source_name,
                "records": len(rows),
                "cities": len(cities),
                "fallback": False,
            },
            "cities": cities,
            "rows": rows,
        }
    finally:
        wb.close()


@st.cache_data(ttl=REFRESH_SECONDS, show_spinner=False)
def load_remote_data(url: str) -> dict[str, Any]:
    content, origin = download_source(url)
    digest = hashlib.sha256(content).hexdigest()[:12]
    data = extract_records_from_xlsx(content, f"{origin} · {digest}")
    return data


@st.cache_data(show_spinner=False)
def load_fallback_data() -> dict[str, Any]:
    return json.loads(FALLBACK_PATH.read_text(encoding="utf-8"))


def load_data() -> tuple[dict[str, Any], str | None]:
    url = secret("MDU_SOURCE_URL", DEFAULT_SOURCE_URL)
    if url:
        try:
            return load_remote_data(url), None
        except Exception as exc:
            fallback = load_fallback_data()
            return fallback, f"Não foi possível atualizar a planilha agora. Exibindo a última base embarcada. Detalhe: {exc}"
    return load_fallback_data(), "A fonte automática ainda não foi configurada. Exibindo a base inicial embarcada no projeto."


def require_password() -> None:
    configured = secret("APP_PASSWORD")
    if not configured:
        return
    if st.session_state.get("auth_ok"):
        return

    st.markdown("<div class='login-title'>Consulta MDU Óptica</div>", unsafe_allow_html=True)
    password = st.text_input("Senha de acesso", type="password", placeholder="Digite a senha")
    if st.button("Entrar", use_container_width=True, type="primary"):
        if hmac.compare_digest(password, configured):
            st.session_state.auth_ok = True
            st.rerun()
        st.error("Senha incorreta.")
    st.stop()


CSS = r"""
<style>
#MainMenu, footer, header { visibility: hidden; }
[data-testid="stAppViewContainer"] {
    background:
        radial-gradient(circle at top right, rgba(21,166,184,.12), transparent 28%),
        linear-gradient(180deg, #f4f7fb 0%, #eef3f8 40%, #eef3f8 100%);
}
[data-testid="stMainBlockContainer"] { max-width: 1220px; padding-top: 1.2rem; padding-bottom: 3rem; }
.block-container { padding-left: 1rem; padding-right: 1rem; }
section[data-testid="stSidebar"] { display:none; }
.hero {
    background: linear-gradient(135deg, #10213d 0%, #16315a 62%, #1a4670 100%);
    color: white;
    border: 1px solid rgba(255,255,255,.08);
    border-bottom: 4px solid #15a6b8;
    border-radius: 26px;
    padding: 28px 28px 22px;
    margin-bottom: 18px;
    box-shadow: 0 20px 45px rgba(16,33,61,.18);
    position: relative;
    overflow: hidden;
}
.hero:before {
    content: "";
    position: absolute;
    right: -60px;
    top: -60px;
    width: 180px;
    height: 180px;
    border-radius: 50%;
    background: rgba(255,255,255,.06);
}
.hero-grid { display:grid; grid-template-columns: minmax(0, 1.35fr) minmax(260px, .9fr); gap: 20px; align-items: start; position:relative; z-index:1; }
.hero-left { display:flex; gap:16px; align-items:flex-start; }
.hero-icon {
    width: 64px; height: 64px; border-radius: 18px;
    display:flex; align-items:center; justify-content:center;
    background: rgba(255,255,255,.09);
    border: 1px solid rgba(255,255,255,.16);
    font-size: 1.9rem; flex: 0 0 auto;
    box-shadow: inset 0 1px 0 rgba(255,255,255,.10);
}
.hero-title { font-size: 2rem; font-weight: 800; letter-spacing: -.03em; line-height: 1.1; margin: 0; }
.hero-sub { color:#c8d6ea; margin-top:4px; font-size:1rem; }
.hero-help { color:#dce7f4; margin-top:14px; font-size:.88rem; line-height:1.5; max-width: 740px; }
.hero-stats {
    display:grid; grid-template-columns: repeat(2, minmax(0,1fr)); gap: 10px;
}
.hero-stat {
    padding: 12px 14px; border-radius: 16px; min-height: 84px;
    background: rgba(255,255,255,.08);
    border: 1px solid rgba(255,255,255,.14);
    backdrop-filter: blur(2px);
}
.hero-stat-label { color:#b7cae3; font-size:.72rem; font-weight:700; text-transform:uppercase; letter-spacing:.04em; }
.hero-stat-value { color:white; font-size:1.05rem; font-weight:800; margin-top:7px; line-height:1.25; }
.source-line { color:#71829b; font-size:.82rem; margin:.15rem .2rem 1rem; }
.kpi-row { display:grid; grid-template-columns: repeat(4, minmax(0,1fr)); gap: 12px; margin-bottom: 18px; }
.kpi-card {
    background: rgba(255,255,255,.86);
    border: 1px solid #d8e1eb;
    border-radius: 18px;
    padding: 14px 16px;
    box-shadow: 0 10px 24px rgba(16,33,61,.06);
}
.kpi-label { color:#69788e; font-size:.74rem; font-weight:800; text-transform:uppercase; letter-spacing:.035em; }
.kpi-value { color:#12233f; font-size:1.15rem; font-weight:800; margin-top:7px; line-height:1.25; }
.search-card {
    background: rgba(255,255,255,.95);
    border:1px solid #d8e1eb;
    border-radius:22px;
    padding:20px 20px 16px;
    box-shadow:0 16px 34px rgba(16,33,61,.08);
    margin-bottom:18px;
}
.search-title { color:#12233f; font-size:1.08rem; font-weight:800; margin-bottom:4px; }
.search-sub { color:#617389; font-size:.86rem; margin-bottom:14px; }
.result-toolbar {
    display:flex; align-items:center; justify-content:space-between; gap:12px; flex-wrap:wrap;
    margin: 6px 0 12px;
}
.result-count {
    color:#12233f; font-size:1rem; font-weight:800;
}
.result-chip {
    display:inline-flex; align-items:center; gap:8px;
    padding: 8px 12px; border-radius:999px; background:#fff; border:1px solid #d8e1eb;
    color:#51627a; font-size:.82rem; font-weight:700;
    box-shadow:0 6px 16px rgba(16,33,61,.05);
}
.result-card {
    background:white; border:1px solid #d8e1eb; border-radius:18px; padding:18px 18px 16px; margin:0 0 12px;
    box-shadow:0 8px 22px rgba(16,33,61,.06);
}
.result-top { display:flex; justify-content:space-between; gap:14px; align-items:flex-start; flex-wrap:wrap; }
.result-address { font-size:1.05rem; font-weight:800; color:#12233f; margin-bottom:5px; line-height:1.35; }
.result-city { color:#5f7188; font-size:.83rem; }
.status-pill {
    display:inline-flex; align-items:center; gap:8px; white-space:nowrap;
    border-radius:999px; padding:8px 12px; font-size:.78rem; font-weight:800; border:1px solid transparent;
}
.status-pill.ok { background:#e9f7f2; color:#0d6b4a; border-color:#bce4d4; }
.status-pill.pending { background:#fff6e6; color:#8a5000; border-color:#efd9a3; }
.status-pill.empty { background:#f4f6f9; color:#66788f; border-color:#e1e8f0; }
.fields { display:grid; grid-template-columns: repeat(3, minmax(0,1fr)); gap:10px; margin-top:14px; }
.field-box {
    background:#f8fbff; border:1px solid #e2eaf3; border-radius:14px; padding:12px; min-height:72px;
}
.field-label { color:#73839a; font-size:.68rem; font-weight:800; text-transform:uppercase; letter-spacing:.04em; }
.field-value { color:#12233f; font-size:.92rem; font-weight:760; margin-top:6px; overflow-wrap:anywhere; line-height:1.35; }
.field-box.status-box { background:#fbfcfe; }
.empty {
    border:1px dashed #bac8d8; border-radius:18px; background:rgba(255,255,255,.78); padding:42px 24px; text-align:center; color:#526178;
}
.login-title { font-size:1.7rem; font-weight:800; color:#10213d; margin:3rem 0 1rem; }
[data-testid="stSelectbox"] label, [data-testid="stTextInput"] label { font-weight: 800 !important; color:#33475f !important; }
[data-testid="stTextInput"] input,
[data-testid="stSelectbox"] div[data-baseweb="select"] > div {
    border-radius:14px !important; min-height:52px; background:#fbfdff;
    border:1px solid #ccd8e6 !important;
}
[data-testid="stTextInput"] input:focus,
[data-testid="stSelectbox"] div[data-baseweb="select"] > div:focus-within {
    border-color:#1769e0 !important; box-shadow:0 0 0 4px rgba(23,105,224,.10) !important;
}
.stButton > button {
    border-radius:14px; min-height:48px; font-weight:800; border:1px solid #cfd9e5;
}
.stButton button[kind="secondary"] { background:white; color:#243a56; }
div[data-testid="stNumberInput"] input {
    border-radius:12px !important;
}
@media (max-width: 980px) {
    .hero-grid { grid-template-columns: 1fr; }
    .kpi-row { grid-template-columns: repeat(2, minmax(0,1fr)); }
    .fields { grid-template-columns: 1fr 1fr; }
}
@media (max-width: 760px) {
    [data-testid="stMainBlockContainer"] { padding-top: .4rem; }
    .hero {
        border-radius: 0 0 22px 22px; margin-left:-1rem; margin-right:-1rem; margin-top:-.5rem; padding: 20px 18px 18px;
    }
    .hero-left { gap:12px; }
    .hero-icon { width:56px; height:56px; font-size:1.6rem; border-radius:16px; }
    .hero-title { font-size: 1.55rem; }
    .hero-sub { font-size:.92rem; }
    .hero-help { font-size:.84rem; }
    .hero-stats { grid-template-columns: 1fr 1fr; }
    .kpi-row { grid-template-columns: 1fr 1fr; gap:10px; }
    .search-card { padding:16px 14px 14px; border-radius:18px; }
    .fields { grid-template-columns: 1fr; }
}
@media (max-width: 430px) {
    .hero-stats, .kpi-row { grid-template-columns: 1fr; }
}
</style>
"""
st.markdown(CSS, unsafe_allow_html=True)
require_password()

st_autorefresh(interval=REFRESH_SECONDS * 1000, key="mdu_auto_refresh")

with st.spinner("Carregando base MDU..."):
    data, load_warning = load_data()

meta = data.get("meta", {})
records = data.get("rows", [])
cities = data.get("cities", [])
base_date = meta.get("sourceDate") or "não informada"
verified_at = meta.get("updatedAt") or "não informado"
source_label = "Google Drive · automático" if not meta.get("fallback") else "Base de contingência"
source_file = str(meta.get("sourceFile") or "")

st.markdown(
    f"""
    <section class="hero">
      <div class="hero-grid">
        <div>
          <div class="hero-left">
            <div class="hero-icon">🔎</div>
            <div>
              <div class="hero-title">Consulta MDU Óptica</div>
              <div class="hero-sub">Consulta operacional · somente leitura</div>
            </div>
          </div>
          <div class="hero-help">
            Pesquise por <strong>endereço</strong> ou <strong>node</strong>, filtre por cidade e consulte rapidamente os 6 campos operacionais da base MDU.
          </div>
        </div>
        <div class="hero-stats">
          <div class="hero-stat"><div class="hero-stat-label">Fonte</div><div class="hero-stat-value">{source_label}</div></div>
          <div class="hero-stat"><div class="hero-stat-label">Base</div><div class="hero-stat-value">{base_date}</div></div>
          <div class="hero-stat"><div class="hero-stat-label">Última verificação</div><div class="hero-stat-value">{verified_at}</div></div>
          <div class="hero-stat"><div class="hero-stat-label">Cobertura</div><div class="hero-stat-value">{len(records):,} registros · {len(cities)} cidades</div></div>
        </div>
      </div>
    </section>
    """.replace(",", "."),
    unsafe_allow_html=True,
)
if source_file:
    st.markdown(f"<div class='source-line'>Fonte: {source_file}</div>", unsafe_allow_html=True)

if load_warning:
    st.warning(load_warning)

st.markdown(
    f"""
    <div class="kpi-row">
      <div class="kpi-card"><div class="kpi-label">Atualização automática</div><div class="kpi-value">a cada {REFRESH_SECONDS // 60} min</div></div>
      <div class="kpi-card"><div class="kpi-label">Campos da consulta</div><div class="kpi-value">Cidade, endereço, node e status</div></div>
      <div class="kpi-card"><div class="kpi-label">Base atual</div><div class="kpi-value">{len(records):,} registros válidos</div></div>
      <div class="kpi-card"><div class="kpi-label">Abrangência</div><div class="kpi-value">{len(cities)} cidades disponíveis</div></div>
    </div>
    """.replace(",", "."),
    unsafe_allow_html=True,
)

if "city" not in st.session_state:
    st.session_state.city = "Todas as cidades"
if "query" not in st.session_state:
    st.session_state.query = ""
if "page" not in st.session_state:
    st.session_state.page = 1

st.markdown('<div class="search-card">', unsafe_allow_html=True)
st.markdown('<div class="search-title">Filtros de consulta</div><div class="search-sub">Selecione uma cidade e pesquise por endereço ou node.</div>', unsafe_allow_html=True)
col_city, col_search = st.columns([1, 1.8])
with col_city:
    city = st.selectbox("Cidade", ["Todas as cidades", *cities], key="city")
with col_search:
    query = st.text_input("Endereço ou node", key="query", placeholder="Ex.: República, ZERO HORA 1811 ou GPONA01")

b1, b2 = st.columns([1, 1])
with b1:
    if st.button("Limpar filtros", use_container_width=True):
        st.session_state.city = "Todas as cidades"
        st.session_state.query = ""
        st.session_state.page = 1
        st.rerun()
with b2:
    if st.button("Verificar atualização", use_container_width=True, type="primary"):
        load_remote_data.clear()
        st.session_state.page = 1
        st.rerun()
st.caption("A busca ignora acentos, diferenças entre Rua/R e espaços extras.")
st.markdown('</div>', unsafe_allow_html=True)

query_norm = normalize_search(query)
if query_norm and len(query_norm) < 2:
    filtered: list[dict[str, str]] = []
    instruction = "Digite ao menos dois caracteres para pesquisar."
elif city == "Todas as cidades" and not query_norm:
    filtered = []
    instruction = "Pesquise por endereço/node ou selecione uma cidade para listar os registros."
else:
    terms = query_norm.split()
    filtered = []
    for item in records:
        if city != "Todas as cidades" and item.get("c") != city:
            continue
        if terms:
            haystack = normalize_search(f"{item.get('a','')} {item.get('n','')}")
            if not all(term in haystack for term in terms):
                continue
        filtered.append(item)
    instruction = ""

if instruction:
    st.markdown(f'<div class="empty">{instruction}</div>', unsafe_allow_html=True)
else:
    total = len(filtered)
    total_pages = max(1, (total + PAGE_SIZE - 1) // PAGE_SIZE)
    if st.session_state.page > total_pages:
        st.session_state.page = 1

    result_city_label = city if city != "Todas as cidades" else "múltiplas cidades"
    st.markdown(
        f"""
        <div class="result-toolbar">
          <div class="result-count">{total:,} resultado(s)</div>
          <div class="result-chip">Cidade: {result_city_label}</div>
        </div>
        """.replace(",", "."),
        unsafe_allow_html=True,
    )

    if total_pages > 1:
        page = st.number_input("Página", min_value=1, max_value=total_pages, value=st.session_state.page, step=1)
        st.session_state.page = int(page)
    else:
        page = 1

    start = (int(page) - 1) * PAGE_SIZE
    end = start + PAGE_SIZE
    for item in filtered[start:end]:
        status = item.get("s") or "Não informado"
        norm_status = normalize_search(status)
        if "CONCLU" in norm_status:
            pill_class = "ok"
        elif status != "Não informado":
            pill_class = "pending"
        else:
            pill_class = "empty"

        st.markdown(
            f"""
            <article class="result-card">
              <div class="result-top">
                <div>
                  <div class="result-address">{item.get('a') or 'Endereço não informado'}</div>
                  <div class="result-city">{item.get('c') or 'Cidade não informada'}</div>
                </div>
                <div class="status-pill {pill_class}">Status RI · {status}</div>
              </div>
              <div class="fields">
                <div class="field-box"><div class="field-label">Node</div><div class="field-value">{item.get('n') or 'Não informado'}</div></div>
                <div class="field-box"><div class="field-label">EPO MDU responsável</div><div class="field-value">{item.get('e') or 'Não informado'}</div></div>
                <div class="field-box"><div class="field-label">Construção RI MDU – fim</div><div class="field-value">{item.get('f') or 'Não informado'}</div></div>
              </div>
            </article>
            """,
            unsafe_allow_html=True,
        )

    if total_pages > 1:
        st.caption(f"Página {int(page)} de {total_pages} · {PAGE_SIZE} registros por página")
